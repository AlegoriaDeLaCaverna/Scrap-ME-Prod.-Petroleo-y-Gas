import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font
import urllib3
import io
import warnings

warnings.filterwarnings('ignore')
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

st.set_page_config(page_title="Monitor de Hidrocarburos", layout="wide", page_icon="🛢️")

# --- BLINDAJE VISUAL ---
st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    </style>
    """, unsafe_allow_html=True)

st.title("🛢️ Monitor Analítico de Hidrocarburos")
st.markdown("Datos extraídos del Ministerio de Energía, estandarizados a métricas financieras corporativas (Kbbl y KBoe).")

diccionario_fuentes = {
    'Petróleo: Producción Total': ("http://datos.energia.gob.ar/dataset/590d1284-fd6d-4686-afd8-b3da5d90a6e9/resource/4cc61040-aa44-440d-a912-91bd6c26b8a7/download/produccin-petrleo-sesco-tight-y-shale-captulo-iv-por-empresa.csv", "Petróleo Total", "m3"),
    'Petróleo: Promedio Diario': ("http://datos.energia.gob.ar/dataset/590d1284-fd6d-4686-afd8-b3da5d90a6e9/resource/2c1f455e-0103-4d51-8f94-a49c939ac0a1/download/produccin-de-petrleo-promedio-diaria-por-empresa.csv", "Petróleo Promedio", "m3/día"),
    'Gas: Producción Total': ("http://datos.energia.gob.ar/dataset/590d1284-fd6d-4686-afd8-b3da5d90a6e9/resource/63129e00-6a96-4d6e-9ce1-9e6c60287e16/download/produccin-gas-sesco-tight-y-shale-captulo-iv-por-empresa.csv", "Gas Total", "Miles de m3"),
    'Gas: Promedio Diario': ("http://datos.energia.gob.ar/dataset/590d1284-fd6d-4686-afd8-b3da5d90a6e9/resource/419094dd-2905-4ac3-9398-e81513013e5e/download/produccin-de-gas-promedio-diaria-por-empresa.csv", "Gas Promedio", "Miles de m3/día")
}

@st.cache_data(show_spinner=False)
def cargar_datos(url):
    try:
        df = pd.read_csv(url, encoding='utf-8-sig', sep=None, engine='python')
    except UnicodeDecodeError:
        df = pd.read_csv(url, encoding='latin-1', sep=None, engine='python')

    df.columns = df.columns.str.lower().str.strip().str.replace('ï»¿', '')

    for col in df.columns:
        if 'empresa' in col or 'operador' in col: df.rename(columns={col: 'empresa'}, inplace=True)
        elif 'anio' in col or 'año' in col or 'year' in col: df.rename(columns={col: 'anio'}, inplace=True)
        elif col in ['mes', 'month']: df.rename(columns={col: 'mes'}, inplace=True)

    cols_base = ['empresa', 'anio', 'mes']
    faltantes = [c for c in cols_base if c not in df.columns]
    if faltantes:
        st.error(f"El Estado modificó las columnas. Faltan: {faltantes}")
        st.stop()

    columnas_prohibidas = ['indice_tiempo', 'id']
    cols_metricas = [c for c in df.columns if c not in cols_base and c not in columnas_prohibidas and df[c].dtype in [np.float64, np.int64]]

    if not cols_metricas:
        st.error("No hay columnas numéricas de volumen en este archivo.")
        st.stop()

    df = df.dropna(subset=cols_base)
    df['anio'] = pd.to_numeric(df['anio'], errors='coerce').fillna(0).astype(int)
    df['mes'] = pd.to_numeric(df['mes'], errors='coerce').fillna(0).astype(int)
    df = df[df['anio'] > 0]
    return df, cols_metricas

if 'num_grupos' not in st.session_state:
    st.session_state.num_grupos = 2
if 'procesar_clicked' not in st.session_state:
    st.session_state.procesar_clicked = False

# --- PANELES DE CONTROL SUPERIORES ---
st.markdown("---")
c_params1, c_params2 = st.columns([1, 2])

with c_params1:
    st.subheader("1. Origen de Datos")
    seleccion_fuente = st.selectbox("Seleccione la dimensión a analizar:", list(diccionario_fuentes.keys()), label_visibility="collapsed")

url_csv, nombre_fluido, unidad_medida = diccionario_fuentes[seleccion_fuente]
es_petroleo = "Petróleo" in nombre_fluido
es_promedio = "Promedio" in nombre_fluido
es_desglosable = "Total" in seleccion_fuente

with st.spinner('⏳ Extrayendo pergaminos del Ministerio...'):
    df, cols_metricas = cargar_datos(url_csv)

anios_disponibles = sorted(df['anio'].unique().tolist())
lista_empresas_reales = sorted(df['empresa'].unique().tolist())

with c_params2:
    st.subheader("2. Alcance Temporal")
    anio_desde, anio_hasta = st.slider("Seleccione el rango de años de extracción:", 
                                       min_value=min(anios_disponibles), 
                                       max_value=max(anios_disponibles), 
                                       value=(min(anios_disponibles), max(anios_disponibles)),
                                       label_visibility="collapsed")

st.markdown("---")
st.subheader("3. Estrategia Corporativa")

col_estrat1, col_estrat2 = st.columns([1, 2])

with col_estrat1:
    modo_analisis = st.radio("Enfoque Analítico:", ["Consolidar (Sumar total)", "Comparar Grupos/Empresas"])
    
grupos_definidos = {}
empresas_elegidas_consolidar = []

with col_estrat2:
    if modo_analisis == "Consolidar (Sumar total)":
        empresas_elegidas_consolidar = st.multiselect("🏭 Seleccione Empresas a consolidar (Deje vacío para analizar Todas):", options=lista_empresas_reales)
        if es_desglosable:
            tipo_grafico = st.radio("📈 Segmentación:", ["Total Consolidado", "Convencional vs No Convencional"], horizontal=True)
        else:
            tipo_grafico = "Total Consolidado"
    else:
        st.caption("Funda personerías jurídicas bajo un mismo nombre para ponerlas a competir.")
        c_btn1, c_btn2, _ = st.columns([1, 1, 2])
        if c_btn1.button("➕ Agregar Competidor"): st.session_state.num_grupos += 1
        if c_btn2.button("➖ Quitar Competidor") and st.session_state.num_grupos > 1: st.session_state.num_grupos -= 1

        for i in range(st.session_state.num_grupos):
            cx1, cx2 = st.columns([1, 2])
            nombre_g = cx1.text_input(f"Competidor {i+1} (Nombre):", key=f"nom_{i}")
            sociedades_g = cx2.multiselect(f"Sociedades de Competidor {i+1}:", options=lista_empresas_reales, key=f"emps_{i}")
            if nombre_g and sociedades_g:
                grupos_definidos[nombre_g] = sociedades_g

st.markdown("<br>", unsafe_allow_html=True)
if st.button("🎯 Procesar Gráfico", type="primary", use_container_width=True):
    st.session_state.procesar_clicked = True

if st.session_state.procesar_clicked:
    df_filtrado_base = df[(df['anio'] >= anio_desde) & (df['anio'] <= anio_hasta)].copy()
    factor_conversion = 6.2898 / 1000

    if es_petroleo:
        unidad_final = "Miles de Barriles/día" if es_promedio else "Miles de Barriles"
        acronimo_final = "Kbbl/d" if es_promedio else "Kbbl"
    else:
        unidad_final = "Miles de BOE/día" if es_promedio else "Miles de BOE"
        acronimo_final = "KBoe/d" if es_promedio else "KBoe"

    # ================= RAMA 1: COMPARACIÓN =================
    if modo_analisis == "Comparar Grupos/Empresas":
        if not grupos_definidos:
            st.warning("Debe definir al menos un competidor.")
            st.stop()
            
        dfs_grupos = []
        for nombre_grupo, sociedades in grupos_definidos.items():
            temp_df = df_filtrado_base[df_filtrado_base['empresa'].isin(sociedades)].copy()
            temp_df['Competidor'] = nombre_grupo
            dfs_grupos.append(temp_df)
            
        df_filtrado = pd.concat(dfs_grupos)
        if df_filtrado.empty:
            st.warning("El pozo está seco para las sociedades seleccionadas en este rango temporal.")
            st.stop()
            
        # Creamos fecha real para Plotly
        df_filtrado['Fecha'] = pd.to_datetime(df_filtrado['anio'].astype(str) + '-' + df_filtrado['mes'].astype(str) + '-01')
        df_filtrado['Periodo_Str'] = df_filtrado['anio'].astype(str) + "-" + df_filtrado['mes'].astype(str).str.zfill(2)
        col_valor = cols_metricas[0]
        
        df_agrupado = df_filtrado.groupby(['Fecha', 'Periodo_Str', 'Competidor'])[col_valor].sum().reset_index()
        df_pivot = df_agrupado.pivot(index=['Fecha', 'Periodo_Str'], columns='Competidor', values=col_valor).reset_index().fillna(0)
        df_pivot.columns.name = None
        
        cols_grupos = [c for c in df_pivot.columns if c not in ['Fecha', 'Periodo_Str']]
        cols_convertidas = []
        
        for grp in cols_grupos:
            col_conv = f"{grp} ({acronimo_final})"
            df_pivot[col_conv] = (df_pivot[grp] * factor_conversion).round(2)
            cols_convertidas.append(col_conv)
            
        cols_to_graph_web = cols_convertidas
        encabezados_excel = ['Período'] + [f"{g} (Orig)" for g in cols_grupos] + cols_convertidas
        nombre_empresa_reporte = " vs ".join(list(grupos_definidos.keys())[:3])
        if len(grupos_definidos) > 3: nombre_empresa_reporte += " y otros"
        df_final = df_pivot.sort_values('Fecha')

    # ================= RAMA 2: CONSOLIDACIÓN =================
    else:
        if not empresas_elegidas_consolidar:
            nombre_empresa_reporte = "(Todas)"
            df_filtrado = df_filtrado_base.copy()
        else:
            df_filtrado = df_filtrado_base[df_filtrado_base['empresa'].isin(empresas_elegidas_consolidar)].copy()
            nombre_empresa_reporte = empresas_elegidas_consolidar[0] if len(empresas_elegidas_consolidar) == 1 else (" + ".join(empresas_elegidas_consolidar) if len(empresas_elegidas_consolidar) <= 3 else "Consorcio Consolidado")
                
        if df_filtrado.empty:
            st.warning("El pozo está seco. No hay registros para esta combinación.")
            st.stop()
            
        df_filtrado['Fecha'] = pd.to_datetime(df_filtrado['anio'].astype(str) + '-' + df_filtrado['mes'].astype(str) + '-01')
        df_filtrado['Periodo_Str'] = df_filtrado['anio'].astype(str) + "-" + df_filtrado['mes'].astype(str).str.zfill(2)
        
        if es_desglosable and 'concepto' in df_filtrado.columns:
            col_valor = cols_metricas[0] 
            df_pivot = df_filtrado.pivot_table(index=['Fecha', 'Periodo_Str'], columns='concepto', values=col_valor, aggfunc='sum').reset_index().fillna(0)
            cols_grafico = [c for c in df_pivot.columns if c not in ['Fecha', 'Periodo_Str']]
        else:
            df_pivot = df_filtrado.groupby(['Fecha', 'Periodo_Str'])[cols_metricas].sum().reset_index()
            cols_grafico = cols_metricas
            
        df_final = df_pivot.sort_values('Fecha')
        nombre_col_nueva = f"{'Promedio' if es_promedio else 'Total'} Consolidado ({acronimo_final})"
        df_final[nombre_col_nueva] = (df_final[cols_grafico].sum(axis=1) * factor_conversion).round(2)
        
        if es_desglosable and 'concepto' in df_filtrado.columns:
            cols_no_conv_raw = [c for c in cols_grafico if any(k in str(c).lower() for k in ['shale', 'tight', 'no convencional'])]
            cols_conv_raw = [c for c in cols_grafico if c not in cols_no_conv_raw]
            
            raw_no_conv = df_final[cols_no_conv_raw].sum(axis=1) if cols_no_conv_raw else 0
            raw_conv = df_final[cols_conv_raw].sum(axis=1) if cols_conv_raw else 0
            
            col_conv_convrt = f"Convencional ({acronimo_final})"
            col_noconv_convrt = f"No Convencional ({acronimo_final})"
            
            df_final[col_conv_convrt] = (raw_conv * factor_conversion).round(2)
            df_final[col_noconv_convrt] = (raw_no_conv * factor_conversion).round(2)
            
            cols_to_graph_web = [col_conv_convrt, col_noconv_convrt] if tipo_grafico == "Convencional vs No Convencional" else [nombre_col_nueva]
            encabezados_excel = ['Período'] + [str(col).replace('_', ' ').title() for col in cols_grafico] + [col_conv_convrt, col_noconv_convrt, nombre_col_nueva]
        else:
            cols_to_graph_web = [nombre_col_nueva]
            encabezados_excel = ['Período'] + [str(col).replace('_', ' ').title() for col in cols_grafico] + [nombre_col_nueva]

   # --- VITRINA GRÁFICA PLOTLY ---
    st.markdown("---")
    st.subheader(f"📊 {nombre_fluido} - {nombre_empresa_reporte}")
    
    df_melted = df_final.melt(id_vars=['Fecha'], value_vars=cols_to_graph_web, var_name='Categoría', value_name='Volumen')
    
    # Formateamos la fecha a string YYYY-MM directamente en el dataframe para el hover
    df_melted['Periodo_Format'] = df_melted['Fecha'].dt.strftime('%Y-%m')

    fig = px.line(df_melted, x='Fecha', y='Volumen', color='Categoría',
                  labels={'Volumen': f'Volumen ({unidad_final})', 'Fecha': ''})
    
    # Aplicamos el hovertemplate para lograr la limpieza institucional
    fig.update_traces(
        line_width=2.5,
        hovertemplate="<br>".join([
            "Periodo: %{customdata[0]}",
            "%{data.name}: %{y:,.2f}"
        ]) + "<extra></extra>", # extra></extra> oculta el recuadro secundario molesto de Plotly
        customdata=df_melted[['Periodo_Format']]
    )
    
    fig.update_layout(
        xaxis=dict(tickformat="%b %Y", showgrid=True, gridcolor='rgba(200, 200, 200, 0.2)'),
        yaxis=dict(showgrid=True, gridcolor='rgba(200, 200, 200, 0.2)'),
        plot_bgcolor='rgba(0,0,0,0)',
        paper_bgcolor='rgba(0,0,0,0)',
        legend_title_text='',
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
        margin=dict(l=0, r=0, t=30, b=0),
        hovermode="x unified" # Hace que la línea vertical cruce el gráfico al pasar el mouse
    )
    st.plotly_chart(fig, use_container_width=True)
    # --- CREACIÓN DEL EXCEL ---
    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {nombre_fluido}"
    
    ws.append(encabezados_excel)
    
    # Preparamos el df para el excel (sin la columna Fecha de datetime)
    df_excel = df_final.drop(columns=['Fecha']).copy()
    
    for r in dataframe_to_rows(df_excel, index=False, header=False):
        ws.append(r)
        
    fila_notas = ws.max_row + 2
    ws.cell(row=fila_notas, column=1, value="NOTAS Y FUENTES METODOLÓGICAS:").font = Font(bold=True, italic=True)
    ws.cell(row=fila_notas + 1, column=1, value="• Base de datos: Producción de Petróleo y Gas (SESCO), Min. de Economía.")
    ws.cell(row=fila_notas + 2, column=1, value="• Conversión: Tabla Pampa Energía.")
    ws.cell(row=fila_notas + 3, column=1, value="• Aclaración: Las fuentes oficiales asignan 100% al operador técnico, difiriendo de los balances (Working Interest).")
    
    if es_petroleo: ws.cell(row=fila_notas + 4, column=1, value="• Cálculo: m3 * 6,2898 / 1.000 = Kbbl.")
    else: ws.cell(row=fila_notas + 4, column=1, value="• Cálculo: Mm3 * 6,2898 / 1.000 = KBoe.")
    
    # --- GRÁFICO EXCEL ---
    chart = LineChart()
    chart.title = f"{nombre_fluido} - {nombre_empresa_reporte}"
    chart.style = 13
    chart.width = 24  
    chart.height = 12 
    chart.x_axis.title = "Período"
    chart.y_axis.title = f"Volumen ({unidad_final})"
    chart.y_axis.numFmt = '#,##0'
    
    if modo_analisis == "Comparar Grupos/Empresas":
        min_c = len(df_excel.columns) - len(cols_convertidas) + 1
        max_c = len(df_excel.columns)
        chart.legend.position = 'b'
    elif es_desglosable and tipo_grafico == "Convencional vs No Convencional":
        min_c = len(df_excel.columns) - 1 
        max_c = len(df_excel.columns)
        chart.legend.position = 'b'
    else:
        min_c = len(df_excel.columns)
        max_c = len(df_excel.columns)
        chart.legend = None 
        
    datos_ref = Reference(ws, min_col=min_c, min_row=1, max_col=max_c, max_row=ws.max_row - 5)
    cat_ref = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row - 5)
    
    chart.add_data(datos_ref, titles_from_data=True, from_rows=False)
    chart.set_categories(cat_ref)
    ws.add_chart(chart, "H2")
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    fluido_archivo = "Petroleo" if "Petróleo" in nombre_fluido else "Gas"
    tipo_archivo = "Comparativa" if modo_analisis == "Comparar Grupos/Empresas" else ("Total" if "Total" in nombre_fluido else "Promedio")
    nombre_empresa_limpio = nombre_empresa_reporte.replace(' ', '_').replace('/', '-')[:30]
        
    nombre_archivo = f"Reporte_{fluido_archivo}_{tipo_archivo}_{nombre_empresa_limpio}.xlsx"
    
    st.markdown("---")
    st.download_button(label="👨‍💻 Descargar Archivo Excel Definitivo", data=buffer, file_name=nombre_archivo, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
